from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/fengyuchen/Desktop/UW/👸Self/OLiYuchen.github.io")
SOURCE = Path("/Users/fengyuchen/.codex/attachments/6897f5c6-e4e2-45a8-8549-4f9077427bdf/pasted-text.txt")
OUT = ROOT / "deliverables" / "Corgi_Insurance_尽调报告_正文10页版.docx"

BLUE = "1F4E79"
DARK = "1F2933"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F5F7"
VERY_LIGHT = "FAFBFC"
GOLD_FILL = "FFF7E6"
RED_FILL = "FDECEC"
GREEN_FILL = "EAF7EE"
BORDER = "D0D7DE"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(int(width * 1440)))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_paragraph_border_bottom(paragraph, color=BLUE, size="12", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def set_run_font(run, name="PingFang SC", east_asia="PingFang SC", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_para(paragraph, before=0, after=6, line=1.12, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_text(doc, text, style=None, after=6, before=0):
    p = doc.add_paragraph(style=style)
    style_para(p, before=before, after=after, line=1.04)
    r = p.add_run(text)
    set_run_font(r, size=9.0, color=DARK)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_para(p, before=8 if level == 1 else 5, after=3 if level > 1 else 4, line=1.0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    r = p.add_run(text)
    size = {1: 12.5, 2: 10.3, 3: 9.2}[level]
    color = BLUE if level in {1, 2} else "31556F"
    set_run_font(r, size=size, color=color, bold=True)
    return p


def add_small_caps(doc, text, color=MUTED):
    p = doc.add_paragraph()
    style_para(p, before=0, after=5, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=7.6, color=color, bold=True)
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE, title_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [7.15])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color="FFFFFF", size="0")
    set_cell_margins(cell, top=100, bottom=100, start=140, end=140)
    p = cell.paragraphs[0]
    style_para(p, after=3, line=1.08)
    r = p.add_run(title)
    set_run_font(r, size=8.8, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    style_para(p2, after=0, line=1.12)
    r2 = p2.add_run(body)
    set_run_font(r2, size=8.3, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, font_size=7.8):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        set_cell_border(hdr[i])
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[i].paragraphs[0]
        style_para(p, after=0, line=1.05)
        r = p.add_run(h)
        set_run_font(r, size=font_size, color=DARK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            set_cell_margins(cells[i], top=70, bottom=70, start=100, end=100)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[i].paragraphs[0]
            style_para(p, after=0, line=1.08)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_code_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [7.15])
    cell = table.cell(0, 0)
    set_cell_shading(cell, VERY_LIGHT)
    set_cell_border(cell, color=BORDER, size="4")
    set_cell_margins(cell, top=75, bottom=75, start=120, end=120)
    p = cell.paragraphs[0]
    style_para(p, after=0, line=1.05)
    for i, line in enumerate(text.strip("\n").splitlines()):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name="PingFang SC", east_asia="PingFang SC", size=7.4, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def classify_line(line):
    if line.startswith("# "):
        return "h1", line[2:].strip()
    if line.startswith("## "):
        return "h2", line[3:].strip()
    if line.startswith("### "):
        return "h3", line[4:].strip()
    return None, line


def parse_blocks(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line.strip() == "---":
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(code)))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1
            headers = [c.strip() for c in table_lines[0].strip("|").split("|")]
            rows = []
            for tl in table_lines[2:]:
                rows.append([c.strip() for c in tl.strip("|").split("|")])
            blocks.append(("table", (headers, rows)))
            continue
        kind, val = classify_line(line)
        if kind:
            blocks.append((kind, val))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(("numbered", items))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(("bullets", items))
            continue
        para = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or nxt.strip() == "---" or nxt.startswith("#") or nxt.startswith("```") or nxt.startswith("|") or nxt.startswith("- ") or re.match(r"^\d+\.\s+", nxt):
                break
            para.append(nxt)
            i += 1
        blocks.append(("para", " ".join(para)))
    return blocks


def clean_inline(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"^#+\s*", "", text)
    text = text.replace('"', '"').replace('"', '"')
    return text.strip()


def configure_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(0.58 if attr in {"top_margin", "bottom_margin"} else 0.62))
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "PingFang SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(9.0)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.04

    for level, size, color in [(1, 12.5, BLUE), (2, 10.3, BLUE), (3, 9.2, "31556F")]:
        st = styles[f"Heading {level}"]
        st.font.name = "PingFang SC"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True

    for sec in doc.sections:
        header = sec.header.paragraphs[0]
        header.text = "Corgi Insurance 尽调报告"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(header.runs[0], size=8.5, color=MUTED)
        footer = sec.footer.paragraphs[0]
        footer.text = "资料截止：2026 年 8 月 11 日"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(footer.runs[0], size=8, color=MUTED)
    return doc


def add_cover(doc):
    add_small_caps(doc, "DUE DILIGENCE REPORT", BLUE)
    title = doc.add_paragraph()
    style_para(title, before=4, after=2, line=1.0)
    r = title.add_run("Corgi Insurance 尽调报告")
    set_run_font(r, size=18, color=DARK, bold=True)
    subtitle = doc.add_paragraph()
    style_para(subtitle, after=6, line=1.0)
    r = subtitle.add_run("专业排版版｜资料截止：2026 年 8 月 11 日")
    set_run_font(r, size=8.8, color=MUTED)
    set_paragraph_border_bottom(subtitle, color=BLUE, size="14", space="8")


def add_navigation(doc):
    add_heading(doc, "阅读导航", 1)
    add_table(
        doc,
        ["章节", "读者应关注的问题"],
        [
            ["一、摘要", "六项主要事实与时间维度限制"],
            ["二至四", "公司概况、创始人、法律实体与承保链条"],
            ["五至七", "保险业务、ETF 业务、财务口径辨析"],
            ["八至十", "融资、诉讼争议、结构性风险"],
            ["附录", "履历、公司章程、Corgi Law、证据缺口、查证路径与术语"],
        ],
        [1.55, 4.75],
        font_size=9.4,
    )
    add_callout(
        doc,
        "核心判断",
        "Corgi 的增长叙事并非单一软件公司叙事，而是保险承保、同集团再保、ETF 发行、诉讼与资本市场估值共同叠加后的结构性故事。阅读时应区分品牌层、法律实体层、风险承担层与收入口径层。",
        fill=GOLD_FILL,
        title_color="7A5A00",
    )


def table_widths(headers, rows):
    n = len(headers)
    if n == 2:
        return [1.65, 5.5]
    if n == 3:
        return [1.35, 3.0, 2.8]
    if n == 4:
        return [1.15, 1.7, 1.8, 2.5]
    if n == 5:
        return [1.05, 1.25, 1.55, 1.6, 1.7]
    return [6.3 / n] * n


def add_content_blocks(doc, blocks):
    for kind, content in blocks:
        if kind == "h1":
            text = clean_inline(content)
            if text in {"目录"}:
                continue
            add_heading(doc, text, 1)
        elif kind == "h2":
            add_heading(doc, clean_inline(content), 2)
        elif kind == "h3":
            add_heading(doc, clean_inline(content), 3)
        elif kind == "para":
            text = clean_inline(content)
            if text.startswith("附录一") or text.startswith("附录二"):
                add_heading(doc, text, 1)
            elif len(text) < 34 and ("：" not in text) and text.endswith("说明"):
                add_heading(doc, text, 3)
            else:
                add_text(doc, text)
        elif kind == "bullets":
            for item in content:
                p = doc.add_paragraph(style="List Bullet")
                style_para(p, after=4, line=1.12)
                r = p.add_run(clean_inline(item))
                set_run_font(r, size=10.2, color=DARK)
        elif kind == "numbered":
            for item in content:
                p = doc.add_paragraph(style="List Number")
                style_para(p, after=4, line=1.12)
                r = p.add_run(clean_inline(item))
                set_run_font(r, size=10.2, color=DARK)
        elif kind == "code":
            add_code_box(doc, content)
        elif kind == "table":
            headers, rows = content
            if headers and all(h.strip("- ") == "" for h in headers):
                continue
            widths = table_widths(headers, rows)
            add_table(doc, [clean_inline(h) for h in headers], [[clean_inline(c) for c in row] for row in rows], widths, font_size=7.2 if len(headers) >= 4 else 7.8)


def enhance_risk_section(doc):
    add_heading(doc, "管理层/投资人快速核查清单", 1)
    rows = [
        ["1", "外部再保险", "要求披露 Corgi Reinsurance 是否实际转分保给第三方、分出比例、合约类型与再保人评级。"],
        ["2", "资本与偿付能力", "要求披露 TRRG 与 Corgi Reinsurance 的实缴资本、盈余、准备金政策与压力测试结果。"],
        ["3", "承保质量", "持续跟踪赔付率、综合成本率、长尾险种 IBNR 准备金、拒赔/争议率。"],
        ["4", "客户合规接受度", "抽样检查企业客户合同是否要求 AM Best A- 以上评级，以及 COI 审核是否会升级。"],
        ["5", "ETF 运营经济性", "核查单一管理费收入是否足以覆盖托管、行政、审计、上市与分销等固定成本。"],
    ]
    add_table(doc, ["序号", "核查项", "所需材料/观察指标"], rows, [0.55, 1.45, 4.3], header_fill=LIGHT_BLUE, font_size=9)


def main():
    raw = SOURCE.read_text(encoding="utf-8")
    start = raw.find("# 一、摘要")
    if start != -1:
        raw = raw[start:]
    doc = configure_doc()
    add_cover(doc)
    blocks = parse_blocks(raw)
    add_content_blocks(doc, blocks)
    doc.core_properties.title = "Corgi Insurance 尽调报告"
    doc.core_properties.subject = "公开资料尽调报告，资料截止 2026 年 8 月 11 日"
    doc.core_properties.author = "OpenAI Codex"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
